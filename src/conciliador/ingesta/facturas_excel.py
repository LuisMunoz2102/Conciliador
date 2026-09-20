"""Lector del catálogo de facturas emitidas.

Acepta tanto el archivo de trabajo del contador como el export crudo de MiAdminXML:
las columnas se localizan por el nombre de su encabezado y no por su posición, porque
MiAdminXML exporta 61 columnas y el archivo de trabajo solo 17.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl

from ..core.dinero import a_centavos
from ..core.modelos import Comprobante, MetodoPago, RolFiscal
from ..core.texto import normalizar

# Cada campo se busca por sus encabezados conocidos, en orden de preferencia.
_COLUMNAS = {
    "estado_sat": ["ESTADO SAT"],
    "tipo": ["TIPO", "TIPOCOMPROBANTE"],
    "fecha_emision": ["FECHA EMISION"],
    "serie": ["SERIE"],
    "folio": ["FOLIO"],
    "uuid": ["UUID"],
    "rfc_receptor": ["RFC RECEPTOR"],
    "nombre_receptor": ["NOMBRE RECEPTOR"],
    "subtotal": ["SUBTOTAL"],
    "iva": ["IVA 16%", "IVA"],
    "ret_iva": ["RETENIDO IVA"],
    "ret_isr": ["RETENIDO ISR"],
    "total": ["TOTAL"],
    "moneda": ["MONEDA"],
    "metodo_pago": ["METODO DE PAGO"],
    "estatus_cobro": ["STATUS", "ESTATUS", "ESTADOPAGO"],
}

_COBRADA = {"COBRADA", "PAGADA", "COBRADO", "SI", "X"}


class CatalogoIlegible(RuntimeError):
    pass


def _mapa_de_columnas(encabezados: list) -> dict[str, int]:
    normalizados = {normalizar(str(h)): i for i, h in enumerate(encabezados) if h}
    mapa: dict[str, int] = {}
    for campo, alias in _COLUMNAS.items():
        for a in alias:
            if a in normalizados:
                mapa[campo] = normalizados[a]
                break
    faltantes = {"uuid", "fecha_emision", "total", "rfc_receptor"} - set(mapa)
    if faltantes:
        raise CatalogoIlegible(
            f"al archivo de facturas le faltan columnas: {sorted(faltantes)}. "
            f"Encabezados encontrados: {sorted(normalizados)}"
        )
    return mapa


def _fecha(valor) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(str(valor).strip()[:19], formato).date()
        except ValueError:
            continue
    raise CatalogoIlegible(f"fecha de emisión ilegible: {valor!r}")


def _texto(fila, mapa: dict[str, int], campo: str) -> str:
    i = mapa.get(campo)
    if i is None or i >= len(fila) or fila[i] is None:
        return ""
    return str(fila[i]).strip()


def _metodo(crudo: str) -> MetodoPago:
    n = normalizar(crudo)
    if n.startswith("PUE"):
        return MetodoPago.PUE
    if n.startswith("PPD"):
        return MetodoPago.PPD
    return MetodoPago.DESCONOCIDO


def leer_facturas(ruta: str | Path, hoja: str | None = None) -> list[Comprobante]:
    """Devuelve las facturas vigentes. Las canceladas se descartan."""
    libro = openpyxl.load_workbook(Path(ruta), data_only=True, read_only=True)
    ws = libro[hoja] if hoja else libro[libro.sheetnames[0]]

    filas = ws.iter_rows(values_only=True)
    try:
        encabezados = list(next(filas))
    except StopIteration:
        raise CatalogoIlegible("el archivo de facturas está vacío") from None

    mapa = _mapa_de_columnas(encabezados)
    comprobantes: list[Comprobante] = []

    for numero, fila in enumerate(filas, start=2):
        uuid = _texto(fila, mapa, "uuid")
        if not uuid:
            continue

        estado_sat = _texto(fila, mapa, "estado_sat") or "Vigente"
        if not normalizar(estado_sat).startswith("VIGENTE"):
            continue

        total = a_centavos(_texto(fila, mapa, "total") or "0")
        ret_iva = a_centavos(_texto(fila, mapa, "ret_iva") or "0")
        ret_isr = a_centavos(_texto(fila, mapa, "ret_isr") or "0")
        cobrada = normalizar(_texto(fila, mapa, "estatus_cobro")) in _COBRADA

        comprobante = Comprobante(
            uuid=uuid,
            rol=RolFiscal.EMITIDA,
            tipo=_texto(fila, mapa, "tipo") or "Factura",
            serie=_texto(fila, mapa, "serie") or None,
            folio=_texto(fila, mapa, "folio") or None,
            fecha_emision=_fecha(fila[mapa["fecha_emision"]]),
            rfc_emisor="",
            rfc_receptor=_texto(fila, mapa, "rfc_receptor"),
            nombre_receptor=_texto(fila, mapa, "nombre_receptor"),
            subtotal_cent=a_centavos(_texto(fila, mapa, "subtotal") or "0"),
            iva_cent=a_centavos(_texto(fila, mapa, "iva") or "0"),
            ret_iva_cent=ret_iva,
            ret_isr_cent=ret_isr,
            total_cent=total,
            moneda=_texto(fila, mapa, "moneda") or "MXN",
            metodo_pago=_metodo(_texto(fila, mapa, "metodo_pago")),
            estado_sat=estado_sat,
            monto_cobrable_cent=total - ret_iva - ret_isr,
            saldo_pendiente_cent=0 if cobrada else total,
            marcada_cobrada=cobrada,
            fila_origen=numero,
        )
        comprobantes.append(comprobante)

    libro.close()
    if not comprobantes:
        raise CatalogoIlegible("no se encontró ninguna factura vigente en el archivo")
    return comprobantes
